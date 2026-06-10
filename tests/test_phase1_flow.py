import contextlib
import copy
import io
import json
import tempfile
import unittest
import zipfile
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from wcca_cli.audit import audit_project
from wcca_cli.cli import main
from wcca_cli.engine import calculate_document, write_calculation_outputs
from wcca_cli.extraction import extract_project
from wcca_cli.input_builder import build_input_draft
from wcca_cli.io import read_json, read_yaml
from wcca_cli.knowledge import search_knowledge, show_knowledge, validate_knowledge_cards, validate_model_assets
from wcca_cli.normalization import normalize_project
from wcca_cli.project import ingest_project, init_project
from wcca_cli.registry import ModelRef, file_sha256, get_model, iter_models
from wcca_cli.reporting import validate_results, write_markdown_report
from wcca_cli.review import approve_project, generate_review_checklist
from wcca_cli.validation import has_errors, validate_input_document


SAMPLE_INPUT = Path("examples/project_sample/work/wcca_input.yaml")


class Phase1FlowTest(unittest.TestCase):
    def test_four_phase1_models_are_registered_and_approved(self):
        models = {(model.model_id, model.version, model.metadata["status"]) for model in iter_models()}
        self.assertEqual(
            models,
            {
                ("resistor_divider", "1.0.0", "approved"),
                ("ldo_power_rail", "1.0.0", "approved"),
                ("comparator_threshold", "1.0.0", "approved"),
                ("rc_delay", "1.0.0", "approved"),
            },
        )

    def test_sample_input_validates_for_release(self):
        document = read_yaml(SAMPLE_INPUT)
        issues = validate_input_document(document, release=True)
        self.assertFalse(has_errors(issues), issues)

    def test_release_validation_rejects_unknown_source_document_id(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        document["circuit_blocks"][0]["inputs"]["vin_min"]["source"]["document_id"] = "REQ-UNKNOWN"

        release_issues = validate_input_document(document, release=True)
        trial_issues = validate_input_document(document, release=False)

        self.assertTrue(has_errors(release_issues), release_issues)
        self.assertTrue(any("Unknown source document_id" in issue["message"] for issue in release_issues))
        self.assertFalse(has_errors(trial_issues), trial_issues)

    def test_release_validation_requires_requirement_source_and_review_status(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        requirement = document["circuit_blocks"][0]["requirements"][0]
        requirement.pop("source")
        requirement.pop("review_status")

        issues = validate_input_document(document, release=True)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Missing requirement source" in issue["message"] for issue in issues))
        self.assertTrue(any("Missing requirement review_status" in issue["message"] for issue in issues))

    def test_release_validation_rejects_unknown_requirement_source_document_id(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        document["circuit_blocks"][0]["requirements"][0]["source"]["document_id"] = "REQ-UNKNOWN"

        issues = validate_input_document(document, release=True)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Unknown requirement source document_id" in issue["message"] for issue in issues))

    def test_release_validation_rejects_circuit_type_not_supported_by_selected_model(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        document["circuit_blocks"][0]["type"] = "ldo_power_rail"

        issues = validate_input_document(document, release=True)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("is not supported by model resistor_divider" in issue["message"] for issue in issues))

    def test_release_validation_requires_circuit_block_type(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        del document["circuit_blocks"][0]["type"]

        issues = validate_input_document(document, release=True)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Missing circuit block type" in issue["message"] for issue in issues))

    def test_sample_calculate_validate_report_outputs(self):
        document = read_yaml(SAMPLE_INPUT)
        calculation = calculate_document(document, input_path=SAMPLE_INPUT)
        self.assertEqual(calculation["status"], "calculated")
        self.assertEqual(len(calculation["results"]), 4)
        self.assertEqual(len(calculation["requirements_coverage"]), 4)
        self.assertGreaterEqual(len(calculation["parameter_sources"]), 20)
        divider_vin_min = next(
            item for item in calculation["parameter_sources"]
            if item["circuit_block_id"] == "DIV_5V_TO_ADC" and item["parameter"] == "vin_min"
        )
        self.assertEqual(divider_vin_min["source_type"], "requirement")
        self.assertEqual(divider_vin_min["document_id"], "REQ-SAMPLE")
        for result in calculation["results"]:
            self.assertIn("summary_result", result)
            self.assertIn("calculation_steps", result)
            self.assertIn("traceability", result)
            self.assertEqual(result["release_status"], "release_candidate")
            self.assertEqual(result["summary_result"]["pass_fail"], "pass")

        issues = validate_results(document, calculation)
        self.assertFalse(has_errors(issues), issues)

        output_dir = Path("examples/project_sample/output/test_calculation")
        report_dir = Path("examples/project_sample/output/test_report")
        write_calculation_outputs(output_dir, calculation)
        self.assertTrue((output_dir / "results.json").exists())
        self.assertTrue((output_dir / "model_trace.json").exists())
        self.assertTrue((output_dir / "wcca_results.xlsx").exists())
        self.assertTrue((output_dir / "risk_items.md").exists())
        self.assertTrue((output_dir / "parameter_sources.csv").exists())
        self.assertTrue((output_dir / "requirements_coverage.csv").exists())
        from openpyxl import load_workbook

        workbook = load_workbook(output_dir / "wcca_results.xlsx", read_only=True)
        self.assertIn("traceability", workbook.sheetnames)
        self.assertIn("risk_items", workbook.sheetnames)
        self.assertIn("parameter_sources", workbook.sheetnames)
        self.assertIn("requirements_coverage", workbook.sheetnames)
        self.assertIn("unit_conversions", workbook.sheetnames)
        self.assertGreater(workbook["parameter_sources"].max_row, 1)
        self.assertGreater(workbook["requirements_coverage"].max_row, 1)
        saved = read_json(output_dir / "results.json")
        self.assertEqual(saved["status"], "calculated")
        report_path = write_markdown_report(report_dir, document, saved)
        self.assertTrue(report_path.exists())

    def test_calculation_normalizes_engineering_units_before_model_execution(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        divider = next(block for block in document["circuit_blocks"] if block["id"] == "DIV_5V_TO_ADC")
        divider["inputs"]["vin_min"]["value"] = 4750
        divider["inputs"]["vin_min"]["unit"] = "mV"
        divider["inputs"]["vin_max"]["value"] = 5250
        divider["inputs"]["vin_max"]["unit"] = "mV"
        divider["inputs"]["r_top"]["nominal"]["value"] = 100
        divider["inputs"]["r_top"]["nominal"]["unit"] = "kohm"
        divider["inputs"]["r_bottom"]["nominal"]["value"] = 10
        divider["inputs"]["r_bottom"]["nominal"]["unit"] = "kohm"

        baseline = calculate_document(read_yaml(SAMPLE_INPUT))
        converted = calculate_document(document)

        self.assertEqual(converted["status"], "calculated")
        self.assertFalse(has_errors(validate_input_document(document, release=True)))
        base_summary = next(item for item in baseline["results"] if item["circuit_block_id"] == "DIV_5V_TO_ADC")["summary_result"]
        converted_result = next(item for item in converted["results"] if item["circuit_block_id"] == "DIV_5V_TO_ADC")
        self.assertEqual(converted_result["summary_result"], base_summary)
        converted_parameters = {
            item["parameter"] for item in converted_result["traceability"]["unit_conversions"]
        }
        self.assertIn("vin_min", converted_parameters)
        self.assertIn("r_top.nominal", converted_parameters)

    def test_validate_results_rejects_tampered_margin(self):
        document = read_yaml(SAMPLE_INPUT)
        calculation = calculate_document(document)
        tampered = copy.deepcopy(calculation)
        tampered["results"][0]["summary_result"]["margin_min"] = 99.0

        issues = validate_results(document, tampered)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("margin_min should be" in issue["message"] for issue in issues))

    def test_validate_results_requires_risk_accepted_parameters_in_risk_items(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        tolerance = document["circuit_blocks"][0]["inputs"]["r_top"]["tolerance"]
        tolerance["review_status"] = "risk_accepted"
        tolerance["risk_reason"] = "Supplier tolerance accepted for sample analysis."
        calculation = calculate_document(document)
        self.assertFalse(has_errors(validate_results(document, calculation)))

        tampered = copy.deepcopy(calculation)
        tampered["risk_items"] = [
            item for item in tampered["risk_items"]
            if not (
                item.get("circuit_block_id") == "DIV_5V_TO_ADC"
                and item.get("parameter") == "r_top.tolerance"
            )
        ]
        issues = validate_results(document, tampered)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Missing parameter_risk" in issue["message"] for issue in issues))

    def test_validate_results_rejects_tampered_requirements_coverage(self):
        document = read_yaml(SAMPLE_INPUT)
        calculation = calculate_document(document)
        tampered = copy.deepcopy(calculation)
        tampered["requirements_coverage"][0]["worst_min"] = 0.0

        issues = validate_results(document, tampered)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Coverage worst_min should be" in issue["message"] for issue in issues))

    def test_validate_results_rejects_tampered_parameter_sources(self):
        document = read_yaml(SAMPLE_INPUT)
        calculation = calculate_document(document)
        tampered = copy.deepcopy(calculation)
        tampered["parameter_sources"][0]["document_id"] = "REQ-BAD"

        issues = validate_results(document, tampered)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Parameter source document_id should be" in issue["message"] for issue in issues))

    def test_trial_model_calculation_is_marked_not_for_release(self):
        document = read_yaml(SAMPLE_INPUT)
        base_model = get_model("resistor_divider", "1.0.0")
        draft_model = replace(base_model, metadata={**base_model.metadata, "status": "draft"})

        def fake_get_model(model_id, version=None):
            if model_id == "resistor_divider":
                return draft_model
            return get_model(model_id, version)

        with patch("wcca_cli.engine.get_model", side_effect=fake_get_model), \
                patch("wcca_cli.engine.validate_model_card_binding", return_value=[]):
            calculation = calculate_document(document)

        self.assertEqual(calculation["status"], "trial_calculated")
        self.assertEqual(calculation["release_status"], "not_for_release")
        draft_result = next(item for item in calculation["results"] if item["model_id"] == "resistor_divider")
        self.assertEqual(draft_result["model_status"], "draft")
        self.assertEqual(draft_result["release_status"], "not_for_release")
        issues = validate_results(document, calculation)
        self.assertTrue(has_errors(issues), issues)

    def test_blocked_model_calculation_is_rejected_before_execution(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        document["circuit_blocks"] = [document["circuit_blocks"][0]]
        base_model = get_model("resistor_divider", "1.0.0")
        blocked_model = replace(base_model, metadata={**base_model.metadata, "status": "blocked"})

        with patch("wcca_cli.engine.get_model", return_value=blocked_model), \
                patch("wcca_cli.engine.validate_model_card_binding", return_value=[]):
            calculation = calculate_document(document)

        self.assertEqual(calculation["status"], "model_preflight_failed")
        self.assertEqual(calculation["results"], [])
        self.assertTrue(has_errors(calculation["issues"]), calculation["issues"])

    def test_model_card_hash_mismatch_blocks_calculation_preflight(self):
        document = copy.deepcopy(read_yaml(SAMPLE_INPUT))
        document["circuit_blocks"] = [document["circuit_blocks"][0]]

        with patch("wcca_cli.engine.validate_model_card_binding", return_value=[{
            "level": "error",
            "path": "wcca_models/resistor_divider/1.0.0/model.yaml",
            "message": "sha256 mismatch: expected stale, got current",
        }]):
            calculation = calculate_document(document)

        self.assertEqual(calculation["status"], "model_preflight_failed")
        self.assertEqual(calculation["results"], [])
        self.assertTrue(has_errors(calculation["issues"]), calculation["issues"])
        self.assertIn("sha256 mismatch", calculation["issues"][0]["message"])


class KnowledgeCardBindingTest(unittest.TestCase):
    def test_validate_knowledge_cards_passes(self):
        issues = validate_knowledge_cards()
        self.assertFalse(has_errors(issues), issues)

    def test_model_asset_validation_reports_missing_five_piece_assets(self):
        model = ModelRef(
            model_id="dummy_model",
            version="1.0.0",
            path=Path("tmp_missing_model_assets"),
            metadata={"model_id": "dummy_model", "version": "1.0.0", "status": "approved", "kb_card_id": "KB-DUMMY"},
        )

        issues = validate_model_assets(model)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("Missing model asset: implementation.py" in issue["message"] for issue in issues))
        self.assertTrue(any("at least 3 example" in issue["message"] for issue in issues))
        self.assertTrue(any("at least one test_*.py" in issue["message"] for issue in issues))

    def test_model_asset_validation_reports_insufficient_scenario_tests(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            model_dir = Path(temp_dir)
            (model_dir / "model.yaml").write_text("model_id: dummy_model\n", encoding="utf-8")
            (model_dir / "implementation.py").write_text("# placeholder\n", encoding="utf-8")
            (model_dir / "formula.md").write_text("# Formula\n", encoding="utf-8")
            (model_dir / "review_record.md").write_text("# Review\n", encoding="utf-8")
            examples_dir = model_dir / "examples"
            examples_dir.mkdir()
            for index in range(3):
                (examples_dir / f"case_{index}.yaml").write_text("project: dummy\n", encoding="utf-8")
            tests_dir = model_dir / "tests"
            tests_dir.mkdir()
            (tests_dir / "test_dummy.py").write_text(
                "def test_only_one_case():\n    assert True\n",
                encoding="utf-8",
            )
            model = ModelRef(
                model_id="dummy_model",
                version="1.0.0",
                path=model_dir,
                metadata={
                    "model_id": "dummy_model",
                    "version": "1.0.0",
                    "status": "approved",
                    "kb_card_id": "KB-DUMMY",
                },
            )

            issues = validate_model_assets(model)

        self.assertTrue(has_errors(issues), issues)
        self.assertTrue(any("at least 10 test cases" in issue["message"] for issue in issues))

    def test_documents_index_points_to_existing_local_metadata(self):
        index = read_yaml("wcca_knowledge_base/index/documents_index.yaml")
        document_ids = {document["document_id"] for document in index["documents"]}
        self.assertIn("DS-RES", document_ids)
        self.assertIn("BOM-SAMPLE", document_ids)
        for document in index["documents"]:
            with self.subTest(document_id=document["document_id"]):
                self.assertTrue(Path(document["path"]).exists(), document["path"])

    def test_model_cards_match_model_metadata_and_hashes(self):
        for model in iter_models():
            kb_card_id = model.metadata["kb_card_id"]
            card_path = Path("wcca_knowledge_base/model_cards") / f"{kb_card_id}.yaml"
            self.assertTrue(card_path.exists(), card_path)
            card = read_yaml(card_path)
            self.assertEqual(card["model_id"], model.model_id)
            self.assertEqual(card["model_version"], model.version)
            self.assertEqual(card["status"], model.metadata["status"])
            self.assertEqual(card["kb_card_id"], kb_card_id)

            for source_file in card["source_files"]:
                path = Path(source_file["path"])
                self.assertTrue(path.exists(), path)
                self.assertEqual(source_file["sha256"], file_sha256(path))

    def test_validate_kb_rejects_model_card_required_parameter_drift(self):
        model = get_model("resistor_divider", "1.0.0")
        card_path = Path("wcca_knowledge_base/model_cards") / f"{model.metadata['kb_card_id']}.yaml"
        original = card_path.read_text(encoding="utf-8")
        try:
            card_path.write_text(original.replace("  - vin_min\n", "", 1), encoding="utf-8")

            issues = validate_knowledge_cards()

            self.assertTrue(has_errors(issues), issues)
            self.assertTrue(any(":required_parameters" in issue["path"] for issue in issues))
        finally:
            card_path.write_text(original, encoding="utf-8")

    def test_get_model_defaults_to_approved_version(self):
        model = get_model("resistor_divider")
        self.assertEqual(model.version, "1.0.0")
        self.assertEqual(model.metadata["status"], "approved")

    def test_cli_model_diff_same_version_has_no_differences(self):
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main([
                "model",
                "diff",
                "resistor_divider",
                "--from-version",
                "1.0.0",
                "--to-version",
                "1.0.0",
            ])
        self.assertEqual(status, 0)
        self.assertIn("no differences", stream.getvalue())

    def test_cli_validate_kb_passes(self):
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["validate-kb"])
        self.assertEqual(status, 0)
        self.assertIn("no issues", stream.getvalue())

    def test_knowledge_search_finds_model_card_rule_review_note_and_document(self):
        model_results = search_knowledge("resistor divider")
        self.assertTrue(any(item["id"] == "KB-MODEL-resistor_divider-1.0.0" for item in model_results))
        exact_model_results = search_knowledge("resistor_divider")
        self.assertEqual([item["id"] for item in exact_model_results], ["KB-MODEL-resistor_divider-1.0.0"])
        exact_document_results = search_knowledge("DS-RES")
        self.assertEqual([item["id"] for item in exact_document_results], ["DS-RES"])
        document_results = search_knowledge("RC0402FR-07100KL")
        self.assertEqual([item["id"] for item in document_results], ["BOM-SAMPLE", "DS-RES"])
        rule_results = search_knowledge("typical worst case")
        self.assertTrue(any(item["id"] == "RULE-TYPICAL-001" for item in rule_results))
        note_results = search_knowledge("review checklist")
        self.assertTrue(any(item["doc_type"] == "review_note" for item in note_results))

    def test_knowledge_show_and_cli_search(self):
        item = show_knowledge("KB-MODEL-resistor_divider-1.0.0")
        self.assertIsNotNone(item)
        self.assertEqual(item["doc_type"], "calculation_model_card")
        document_item = show_knowledge("DS-LDO")
        self.assertIsNotNone(document_item)
        self.assertEqual(document_item["doc_type"], "datasheet")

        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["kb", "search", "ldo regulator"])
        self.assertEqual(status, 0)
        self.assertIn("KB-MODEL-ldo_power_rail-1.0.0", stream.getvalue())

        show_stream = io.StringIO()
        with contextlib.redirect_stdout(show_stream):
            show_status = main(["kb", "show", "RULE-SOURCE-001"])
        self.assertEqual(show_status, 0)
        self.assertIn("Release calculations require source", show_stream.getvalue())

        document_stream = io.StringIO()
        with contextlib.redirect_stdout(document_stream):
            document_status = main(["kb", "show", "DS-LDO"])
        self.assertEqual(document_status, 0)
        self.assertIn("Sample 3.3V LDO Datasheet", document_stream.getvalue())


class ProjectBootstrapTest(unittest.TestCase):
    def test_init_and_ingest_project(self):
        project_dir = Path("examples/project_init_smoke")
        init_project(project_dir, "InitSmoke")
        self.assertTrue((project_dir / "work" / "wcca_input.yaml").exists())
        bom_dir = project_dir / "input" / "bom"
        bom_dir.mkdir(parents=True, exist_ok=True)
        (bom_dir / "bom.csv").write_text("ref,part\nR1,10k\n", encoding="utf-8")
        index = ingest_project(project_dir)
        self.assertEqual(len(index["documents"]), 1)
        self.assertEqual(index["documents"][0]["category"], "bom")
        self.assertTrue((project_dir / "work" / "input_index.yaml").exists())

    def test_sample_project_extracts_bom_components_and_parameter_draft(self):
        extracted = extract_project("examples/project_sample")
        self.assertEqual(extracted["summary"]["component_count"], 8)
        self.assertEqual(extracted["summary"]["resistor_count"], 5)
        self.assertEqual(extracted["summary"]["capacitor_count"], 1)
        components_path = Path("examples/project_sample/work/extracted/components.yaml")
        draft_path = Path("examples/project_sample/work/extracted/parameter_draft.yaml")
        self.assertTrue(components_path.exists())
        self.assertTrue(draft_path.exists())
        components = read_yaml(components_path)
        r101 = next(item for item in components["components"] if item["reference"] == "R101")
        self.assertEqual(r101["nominal"]["value"], 100000.0)
        self.assertEqual(r101["nominal"]["unit"], "ohm")
        self.assertEqual(r101["review_status"], "rule_extracted")
        netlist = read_yaml("examples/project_sample/work/extracted/netlist.yaml")
        self.assertEqual(netlist["summary"]["connection_count"], 16)
        self.assertTrue(any(item["candidate_type"] == "rc_pair" for item in netlist["topology_hints"]))
        requirements = read_yaml("examples/project_sample/work/extracted/requirements.yaml")
        self.assertEqual(requirements["summary"]["requirement_count"], 4)
        self.assertEqual(requirements["summary"]["mapped_requirement_count"], 4)

    def test_cli_extract_sample_project(self):
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["extract", "--project", "examples/project_sample"])
        self.assertEqual(status, 0)
        self.assertIn("component_count=8", stream.getvalue())

    def test_sample_project_normalizes_parameters_and_model_candidates(self):
        extract_project("examples/project_sample")
        normalized = normalize_project("examples/project_sample")
        summary = normalized["parameters"]["summary"]
        self.assertEqual(summary["parameter_count"], 12)
        self.assertEqual(summary["trial_only_count"], 12)
        self.assertEqual(summary["release_ready_count"], 0)
        parameters_path = Path("examples/project_sample/work/normalized/parameters.yaml")
        candidates_path = Path("examples/project_sample/work/normalized/model_candidates.yaml")
        self.assertTrue(parameters_path.exists())
        self.assertTrue(candidates_path.exists())
        parameters = read_yaml(parameters_path)
        self.assertIn("R101.nominal", parameters["parameter_index"])
        r101 = next(item for item in parameters["parameters"] if item["parameter_id"] == "R101.nominal")
        self.assertEqual(r101["unit"], "ohm")
        self.assertEqual(r101["issues"], ["requires_engineer_confirmation"])
        candidate_model_ids = {item["model_id"] for item in normalized["model_candidates"]["candidates"]}
        self.assertIn("resistor_divider", candidate_model_ids)
        self.assertIn("comparator_threshold", candidate_model_ids)
        self.assertIn("ldo_power_rail", candidate_model_ids)
        self.assertIn("rc_delay", candidate_model_ids)
        rc_candidate = next(item for item in normalized["model_candidates"]["candidates"] if item["model_id"] == "rc_delay")
        self.assertEqual(set(rc_candidate["component_refs"]), {"R301", "C301"})
        self.assertIn("topology_hint", rc_candidate["evidence"])

    def test_cli_normalize_sample_project(self):
        extract_project("examples/project_sample")
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["normalize", "--project", "examples/project_sample"])
        self.assertEqual(status, 0)
        self.assertIn("parameter_count=12", stream.getvalue())

    def test_sample_project_builds_wcca_input_draft(self):
        extract_project("examples/project_sample")
        normalize_project("examples/project_sample")
        draft = build_input_draft("examples/project_sample")
        self.assertEqual(len(draft["circuit_blocks"]), 4)
        draft_path = Path("examples/project_sample/work/wcca_input_draft.yaml")
        self.assertTrue(draft_path.exists())
        divider = next(block for block in draft["circuit_blocks"] if block["type"] == "resistor_divider")
        self.assertEqual(divider["selected_model"]["selection_status"], "requires_engineer_confirmation")
        self.assertEqual(divider["requirements"][0]["id"], "REQ-DRAFT-DIV")
        self.assertEqual(divider["requirements"][0]["min_value"], 0.415)
        self.assertEqual(divider["requirements"][0]["review_status"], "rule_extracted")
        self.assertEqual(divider["inputs"]["r_top"]["nominal"]["review_status"], "rule_extracted")
        self.assertEqual(divider["inputs"]["vin_min"]["review_status"], "missing")
        rc_block = next(block for block in draft["circuit_blocks"] if block["type"] == "rc_delay")
        self.assertEqual(rc_block["inputs"]["r"]["nominal"]["review_status"], "rule_extracted")
        self.assertEqual(rc_block["inputs"]["c"]["nominal"]["review_status"], "rule_extracted")
        issues = validate_input_document(draft, release=True)
        self.assertTrue(has_errors(issues), issues)

    def test_cli_input_build_sample_project(self):
        extract_project("examples/project_sample")
        normalize_project("examples/project_sample")
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["input-build", "--project", "examples/project_sample"])
        self.assertEqual(status, 0)
        self.assertIn("draft_blocks=4", stream.getvalue())


class ReviewWorkflowTest(unittest.TestCase):
    def test_review_checklist_for_draft_has_blockers(self):
        extract_project("examples/project_sample")
        normalize_project("examples/project_sample")
        build_input_draft("examples/project_sample")
        checklist = generate_review_checklist(
            "examples/project_sample/work/wcca_input_draft.yaml",
            "examples/project_sample/review/comments",
        )
        self.assertGreater(checklist["summary"]["blocking_issue_count"], 0)
        self.assertTrue(Path("examples/project_sample/review/comments/review_checklist.md").exists())

    def test_review_checklist_and_approval_for_release_input(self):
        checklist = generate_review_checklist(
            "examples/project_sample/work/wcca_input.yaml",
            "examples/project_sample/review/comments",
            "examples/project_sample/output/calculation/results.json",
        )
        self.assertEqual(checklist["summary"]["blocking_issue_count"], 0)
        record = approve_project(
            "examples/project_sample/work/wcca_input.yaml",
            "examples/project_sample/output/calculation/results.json",
            "examples/project_sample/review/approvals",
            reviewer="unit_test_reviewer",
            decision="approved",
            comment="Unit test approval.",
        )
        self.assertEqual(record["validation_status"], "passed")
        self.assertTrue(Path("examples/project_sample/review/approvals/approval_record.json").exists())

    def test_cli_review_and_approve(self):
        review_stream = io.StringIO()
        with contextlib.redirect_stdout(review_stream):
            review_status = main([
                "review-checklist",
                "--input",
                "examples/project_sample/work/wcca_input.yaml",
                "--results",
                "examples/project_sample/output/calculation/results.json",
                "--output",
                "examples/project_sample/review/comments",
            ])
        self.assertEqual(review_status, 0)
        self.assertIn("blocking_issues=0", review_stream.getvalue())

        approve_stream = io.StringIO()
        with contextlib.redirect_stdout(approve_stream):
            approve_status = main([
                "approve",
                "--input",
                "examples/project_sample/work/wcca_input.yaml",
                "--results",
                "examples/project_sample/output/calculation/results.json",
                "--output",
                "examples/project_sample/review/approvals",
                "--reviewer",
                "cli_test_reviewer",
                "--decision",
                "approved",
            ])
        self.assertEqual(approve_status, 0)
        self.assertIn("validation_status=passed", approve_stream.getvalue())

    def test_cli_approve_blocks_unfinished_draft(self):
        extract_project("examples/project_sample")
        normalize_project("examples/project_sample")
        build_input_draft("examples/project_sample")
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main([
                "approve",
                "--input",
                "examples/project_sample/work/wcca_input_draft.yaml",
                "--results",
                "examples/project_sample/output/calculation/results.json",
                "--output",
                "examples/project_sample/review/approvals",
                "--reviewer",
                "cli_test_reviewer",
                "--decision",
                "approved",
            ])
        self.assertEqual(status, 1)
        self.assertIn("Cannot approve project", stream.getvalue())


class AuditWorkflowTest(unittest.TestCase):
    def test_project_audit_passes_for_release_package(self):
        approve_project(
            "examples/project_sample/work/wcca_input.yaml",
            "examples/project_sample/output/calculation/results.json",
            "examples/project_sample/review/approvals",
            reviewer="audit_test_reviewer",
            decision="approved",
        )
        main(["package", "--project", "examples/project_sample", "--output", "examples/project_sample/release/wcca_package.zip"])
        with zipfile.ZipFile("examples/project_sample/release/wcca_package.zip") as archive:
            manifest = json.loads(archive.read("archive_manifest.json").decode("utf-8"))
        result_entry = next(item for item in manifest["files"] if item["path"] == "output/calculation/results.json")
        self.assertIn("sha256", result_entry)
        result = audit_project("examples/project_sample")
        self.assertEqual(result["status"], "passed", result["issues"])

    def test_project_audit_fails_on_bad_approval_hash(self):
        approval_path = Path("examples/project_sample/review/approvals/approval_record.json")
        original = read_json(approval_path)
        tampered = dict(original)
        tampered["results_sha256"] = "bad"
        import json

        approval_path.write_text(json.dumps(tampered, indent=2), encoding="utf-8")
        try:
            result = audit_project("examples/project_sample")
            self.assertEqual(result["status"], "failed")
            self.assertTrue(any("Approval results hash" in issue["message"] for issue in result["issues"]))
        finally:
            approval_path.write_text(json.dumps(original, indent=2), encoding="utf-8")

    def test_project_audit_fails_on_package_manifest_hash_mismatch(self):
        approve_project(
            "examples/project_sample/work/wcca_input.yaml",
            "examples/project_sample/output/calculation/results.json",
            "examples/project_sample/review/approvals",
            reviewer="audit_hash_reviewer",
            decision="approved",
        )
        package_path = Path("examples/project_sample/release/wcca_package.zip")
        main(["package", "--project", "examples/project_sample", "--output", str(package_path)])
        rebuilt_path = package_path.with_suffix(".tampered.zip")
        with zipfile.ZipFile(package_path, "r") as source, zipfile.ZipFile(rebuilt_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "output/calculation/results.json":
                    data = data.replace(b'"status": "calculated"', b'"status": "tampered"')
                target.writestr(info, data)
        original_bytes = package_path.read_bytes()
        try:
            package_path.write_bytes(rebuilt_path.read_bytes())
            result = audit_project("examples/project_sample")
            self.assertEqual(result["status"], "failed")
            self.assertTrue(any("Package hash mismatch" in issue["message"] for issue in result["issues"]))
        finally:
            package_path.write_bytes(original_bytes)
            rebuilt_path.unlink(missing_ok=True)

    def test_project_audit_fails_on_coverage_csv_mismatch(self):
        coverage_path = Path("examples/project_sample/output/calculation/requirements_coverage.csv")
        original = coverage_path.read_text(encoding="utf-8")
        try:
            coverage_path.write_text(original.replace("REQ-DIV-001", "REQ-DIV-BAD", 1), encoding="utf-8")
            result = audit_project("examples/project_sample")
            self.assertEqual(result["status"], "failed")
            self.assertTrue(any("Coverage CSV missing row" in issue["message"] for issue in result["issues"]))
        finally:
            coverage_path.write_text(original, encoding="utf-8")

    def test_project_audit_fails_on_parameter_sources_csv_mismatch(self):
        sources_path = Path("examples/project_sample/output/calculation/parameter_sources.csv")
        original = sources_path.read_text(encoding="utf-8")
        try:
            sources_path.write_text(original.replace("REQ-SAMPLE", "REQ-BAD", 1), encoding="utf-8")
            result = audit_project("examples/project_sample")
            self.assertEqual(result["status"], "failed")
            self.assertTrue(any("Parameter sources CSV" in issue["message"] for issue in result["issues"]))
        finally:
            sources_path.write_text(original, encoding="utf-8")

    def test_project_audit_fails_on_model_trace_results_mismatch(self):
        trace_path = Path("examples/project_sample/output/calculation/model_trace.json")
        original = read_json(trace_path)
        tampered = copy.deepcopy(original)
        tampered["results"][0]["release_status"] = "not_for_release"
        trace_path.write_text(json.dumps(tampered, indent=2), encoding="utf-8")
        try:
            result = audit_project("examples/project_sample")
            self.assertEqual(result["status"], "failed")
            self.assertTrue(any("Trace release_status should be" in issue["message"] for issue in result["issues"]))
        finally:
            trace_path.write_text(json.dumps(original, indent=2), encoding="utf-8")

    def test_cli_audit_passes(self):
        approve_project(
            "examples/project_sample/work/wcca_input.yaml",
            "examples/project_sample/output/calculation/results.json",
            "examples/project_sample/review/approvals",
            reviewer="audit_cli_reviewer",
            decision="approved",
        )
        main(["package", "--project", "examples/project_sample", "--output", "examples/project_sample/release/wcca_package.zip"])
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            status = main(["audit", "--project", "examples/project_sample"])
        self.assertEqual(status, 0)
        self.assertIn("audit_status=passed", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
