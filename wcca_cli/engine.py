from __future__ import annotations

from datetime import datetime, timezone
import csv
from pathlib import Path
from typing import Any

from . import __version__
from .io import write_json
from .knowledge import validate_model_card_binding
from .registry import file_sha256, get_model, load_implementation
from .units import normalize_document_units
from .validation import expected_units_for_model, has_errors, validate_input_document


TRIAL_ONLY_MODEL_STATUSES = {"draft", "reviewing"}
BLOCKED_MODEL_STATUSES = {"blocked"}
DEPRECATED_MODEL_STATUSES = {"deprecated"}


def _now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def calculate_document(document: dict[str, Any], input_path: str | Path | None = None) -> dict[str, Any]:
    issues = validate_input_document(document, release=False)
    if has_errors(issues):
        return {"status": "invalid_input", "issues": issues, "results": []}
    expected_units_by_block = _expected_units_by_block(document)
    calculation_document, unit_conversions = normalize_document_units(document, expected_units_by_block)

    results: list[dict[str, Any]] = []
    lifecycle_issues: list[dict[str, Any]] = []
    model_binding_issues: list[dict[str, Any]] = []
    for block in calculation_document.get("circuit_blocks", []):
        selected = block["selected_model"]
        model = get_model(selected["model_id"], selected.get("model_version"))
        lifecycle_issue = _model_lifecycle_issue(block.get("id"), model.metadata.get("status"))
        if lifecycle_issue:
            lifecycle_issues.append(lifecycle_issue)
        for binding_issue in validate_model_card_binding(model):
            copied = dict(binding_issue)
            copied["path"] = f"circuit_blocks[{block.get('id')}].selected_model.{copied.get('path')}"
            model_binding_issues.append(copied)

    preflight_issues = lifecycle_issues + model_binding_issues
    if has_errors(preflight_issues):
        return {"status": "model_preflight_failed", "issues": issues + preflight_issues, "results": []}

    for block in calculation_document.get("circuit_blocks", []):
        selected = block["selected_model"]
        model = get_model(selected["model_id"], selected.get("model_version"))
        module = load_implementation(model)
        calculated = module.calculate(block)
        model_status = model.metadata.get("status")
        traceability = calculated.setdefault("traceability", {})
        traceability.update({
            "model_hash": f"sha256:{file_sha256(model.path / 'model.yaml')}",
            "implementation_hash": f"sha256:{file_sha256(model.path / 'implementation.py')}",
            "input_version": f"sha256:{file_sha256(input_path)}" if input_path else None,
            "calculation_tool_version": __version__,
            "calculated_at": _now_iso(),
            "unit_conversions": [
                item for item in unit_conversions if item.get("circuit_block_id") == block.get("id")
            ],
        })
        calculated["circuit_block_id"] = block.get("id")
        calculated["model_id"] = model.model_id
        calculated["model_version"] = model.version
        calculated["model_status"] = model_status
        calculated["release_status"] = _release_status_for_model(model_status)
        results.append(calculated)
    return {
        "status": "calculated" if not preflight_issues else "trial_calculated",
        "release_status": "not_for_release" if lifecycle_issues else "release_candidate",
        "issues": issues + preflight_issues,
        "results": results,
        "unit_conversions": unit_conversions,
        "risk_items": _collect_risk_items(calculation_document, results),
        "parameter_sources": _build_parameter_sources(calculation_document),
        "requirements_coverage": _build_requirements_coverage(calculation_document, results),
    }


def write_calculation_outputs(output_dir: str | Path, calculation: dict[str, Any]) -> None:
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    write_json(target / "results.json", calculation)
    write_json(target / "model_trace.json", {
        "status": calculation.get("status"),
        "results": [
            {
                "circuit_block_id": item.get("circuit_block_id"),
                "model_id": item.get("model_id"),
                "model_version": item.get("model_version"),
                "model_status": item.get("model_status"),
                "release_status": item.get("release_status"),
                "traceability": item.get("traceability", {}),
            }
            for item in calculation.get("results", [])
        ],
    })
    _write_excel(target / "wcca_results.xlsx", calculation)
    _write_risk_items(target / "risk_items.md", calculation.get("risk_items", []))
    _write_parameter_sources(target / "parameter_sources.csv", calculation.get("parameter_sources", []))
    _write_requirements_coverage(target / "requirements_coverage.csv", calculation.get("requirements_coverage", []))


def _collect_risk_items(document: dict[str, Any], results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    risks: list[dict[str, Any]] = []
    result_by_block = {item.get("circuit_block_id"): item for item in results}
    for block in document.get("circuit_blocks", []):
        block_id = block.get("id")
        result = result_by_block.get(block_id, {})
        summary = result.get("summary_result", {})
        if summary.get("pass_fail") == "fail":
            risks.append({
                "circuit_block_id": block_id,
                "risk_type": "calculation_fail",
                "description": "Worst-case result does not satisfy requirement limits.",
            })
        for side in ("margin_min", "margin_max"):
            margin = summary.get(side)
            if isinstance(margin, (int, float)) and margin < 5:
                risks.append({
                    "circuit_block_id": block_id,
                    "risk_type": "low_margin",
                    "description": f"{side} is below 5%: {margin}",
                })
        _collect_parameter_risks(block.get("inputs", {}), block_id, "", risks)
    return risks


def _collect_parameter_risks(node: Any, block_id: str | None, path: str, risks: list[dict[str, Any]]) -> None:
    if isinstance(node, dict) and "value" in node:
        review_status = node.get("review_status")
        parameter_type = node.get("parameter_type")
        if review_status == "risk_accepted" or parameter_type == "typical":
            risks.append({
                "circuit_block_id": block_id,
                "risk_type": "parameter_risk",
                "parameter": path,
                "description": node.get("risk_reason", f"{path} is {parameter_type or review_status}."),
            })
        return
    if isinstance(node, dict):
        for key, value in node.items():
            child_path = f"{path}.{key}" if path else key
            _collect_parameter_risks(value, block_id, child_path, risks)


def _build_requirements_coverage(document: dict[str, Any], results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result_by_block = {item.get("circuit_block_id"): item for item in results}
    rows: list[dict[str, Any]] = []
    for block in document.get("circuit_blocks", []):
        result = result_by_block.get(block.get("id"), {})
        summary = result.get("summary_result", {})
        selected = block.get("selected_model", {})
        for req in block.get("requirements", []):
            rows.append({
                "requirement_id": req.get("id"),
                "circuit_block_id": block.get("id"),
                "model_id": selected.get("model_id"),
                "model_version": selected.get("model_version"),
                "requirement_min": req.get("min_value"),
                "requirement_max": req.get("max_value"),
                "worst_min": summary.get("worst_min"),
                "worst_max": summary.get("worst_max"),
                "pass_fail": summary.get("pass_fail"),
            })
    return rows


def _build_parameter_sources(document: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for block in document.get("circuit_blocks", []):
        _collect_parameter_sources(block.get("inputs", {}), block, "", rows)
    return rows


def _collect_parameter_sources(
    node: Any,
    block: dict[str, Any],
    path: str,
    rows: list[dict[str, Any]],
) -> None:
    if isinstance(node, dict) and "value" in node:
        source = node.get("source") or {}
        selected = block.get("selected_model", {})
        rows.append({
            "circuit_block_id": block.get("id"),
            "circuit_block_name": block.get("name"),
            "model_id": selected.get("model_id"),
            "model_version": selected.get("model_version"),
            "parameter": path,
            "value": node.get("value"),
            "unit": node.get("unit"),
            "parameter_type": node.get("parameter_type"),
            "source_type": source.get("type") if isinstance(source, dict) else None,
            "document_id": source.get("document_id") if isinstance(source, dict) else None,
            "source_detail": _source_detail(source),
            "review_status": node.get("review_status"),
            "risk_reason": node.get("risk_reason"),
        })
        return
    if isinstance(node, dict):
        for key, value in node.items():
            child_path = f"{path}.{key}" if path else key
            _collect_parameter_sources(value, block, child_path, rows)


def _write_risk_items(path: Path, risks: list[dict[str, Any]]) -> None:
    lines = ["# Risk Items", ""]
    if not risks:
        lines.append("No risk items detected.")
    else:
        for risk in risks:
            lines.append(
                f"- {risk.get('circuit_block_id')}: {risk.get('risk_type')} - {risk.get('description')}"
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_requirements_coverage(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = [
        "requirement_id",
        "circuit_block_id",
        "model_id",
        "model_version",
        "requirement_min",
        "requirement_max",
        "worst_min",
        "worst_max",
        "pass_fail",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_parameter_sources(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = [
        "circuit_block_id",
        "circuit_block_name",
        "model_id",
        "model_version",
        "parameter",
        "value",
        "unit",
        "parameter_type",
        "source_type",
        "document_id",
        "source_detail",
        "review_status",
        "risk_reason",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _source_detail(source: Any) -> str:
    if not isinstance(source, dict):
        return ""
    details = [
        f"{key}={value}"
        for key, value in source.items()
        if key not in {"type", "document_id"} and value not in (None, "")
    ]
    return "; ".join(details)


def _write_excel(path: Path, calculation: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append([
        "circuit_block_id",
        "model_id",
        "model_version",
        "nominal",
        "worst_min",
        "worst_max",
        "requirement_min",
        "requirement_max",
        "margin_min",
        "margin_max",
        "pass_fail",
    ])
    for item in calculation.get("results", []):
        summary = item.get("summary_result", {})
        ws.append([
            item.get("circuit_block_id"),
            item.get("model_id"),
            item.get("model_version"),
            summary.get("nominal"),
            summary.get("worst_min"),
            summary.get("worst_max"),
            summary.get("requirement_min"),
            summary.get("requirement_max"),
            summary.get("margin_min"),
            summary.get("margin_max"),
            summary.get("pass_fail"),
        ])

    steps = wb.create_sheet("calculation_steps")
    steps.append(["circuit_block_id", "step", "formula", "result", "unit"])
    for item in calculation.get("results", []):
        for step in item.get("calculation_steps", []):
            steps.append([
                item.get("circuit_block_id"),
                step.get("step"),
                step.get("formula"),
                step.get("result"),
                step.get("unit"),
            ])

    traceability_rows = []
    for item in calculation.get("results", []):
        traceability = item.get("traceability", {})
        traceability_rows.append({
            "circuit_block_id": item.get("circuit_block_id"),
            "model_id": item.get("model_id"),
            "model_version": item.get("model_version"),
            "model_status": item.get("model_status"),
            "release_status": item.get("release_status"),
            "model_hash": traceability.get("model_hash"),
            "implementation_hash": traceability.get("implementation_hash"),
            "input_version": traceability.get("input_version"),
            "calculation_tool_version": traceability.get("calculation_tool_version"),
            "calculated_at": traceability.get("calculated_at"),
        })
    _append_dict_sheet(wb, "traceability", traceability_rows, [
        "circuit_block_id",
        "model_id",
        "model_version",
        "model_status",
        "release_status",
        "model_hash",
        "implementation_hash",
        "input_version",
        "calculation_tool_version",
        "calculated_at",
    ])

    _append_dict_sheet(wb, "risk_items", calculation.get("risk_items", []), [
        "circuit_block_id",
        "risk_type",
        "parameter",
        "description",
    ])
    _append_dict_sheet(wb, "parameter_sources", calculation.get("parameter_sources", []), [
        "circuit_block_id",
        "circuit_block_name",
        "model_id",
        "model_version",
        "parameter",
        "value",
        "unit",
        "parameter_type",
        "source_type",
        "document_id",
        "source_detail",
        "review_status",
        "risk_reason",
    ])
    _append_dict_sheet(wb, "requirements_coverage", calculation.get("requirements_coverage", []), [
        "requirement_id",
        "circuit_block_id",
        "model_id",
        "model_version",
        "requirement_min",
        "requirement_max",
        "worst_min",
        "worst_max",
        "pass_fail",
    ])
    _append_dict_sheet(wb, "unit_conversions", calculation.get("unit_conversions", []), [
        "circuit_block_id",
        "parameter",
        "from_unit",
        "to_unit",
    ])
    wb.save(path)


def _append_dict_sheet(wb: Any, title: str, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    sheet = wb.create_sheet(title)
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(field) for field in fieldnames])


def _model_lifecycle_issue(block_id: str | None, status: str | None) -> dict[str, Any] | None:
    path = f"circuit_blocks[{block_id}].selected_model"
    if status in BLOCKED_MODEL_STATUSES:
        return {
            "level": "error",
            "path": path,
            "message": f"Model status {status} is blocked and cannot be calculated",
        }
    if status in TRIAL_ONLY_MODEL_STATUSES:
        return {
            "level": "warning",
            "path": path,
            "message": f"Model status {status} is trial-only; results are not_for_release",
        }
    if status in DEPRECATED_MODEL_STATUSES:
        return {
            "level": "warning",
            "path": path,
            "message": "Model status deprecated; use only for historical reproduction unless approved",
        }
    return None


def _release_status_for_model(status: str | None) -> str:
    if status in TRIAL_ONLY_MODEL_STATUSES:
        return "not_for_release"
    if status in DEPRECATED_MODEL_STATUSES:
        return "deprecated_model"
    return "release_candidate"


def _expected_units_by_block(document: dict[str, Any]) -> dict[str | None, dict[str, str]]:
    expected: dict[str | None, dict[str, str]] = {}
    for block in document.get("circuit_blocks", []):
        selected = block.get("selected_model", {})
        try:
            model = get_model(selected.get("model_id"), selected.get("model_version"))
        except KeyError:
            continue
        expected[block.get("id")] = expected_units_for_model(model.metadata)
    return expected
