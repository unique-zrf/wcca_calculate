from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .reporting import write_markdown_report


REPORT_EXPORT_FILES = [
    "output/report/wcca_report.docx",
    "output/report/wcca_report.pdf",
]


def export_report_documents(
    project_dir: str | Path,
    input_document: dict[str, Any],
    results_document: dict[str, Any],
    output_dir: str | Path,
) -> dict[str, str]:
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    markdown_path = write_markdown_report(target, input_document, results_document)
    docx_path = target / "wcca_report.docx"
    pdf_path = target / "wcca_report.pdf"
    _write_docx(docx_path, markdown_path.read_text(encoding="utf-8"))
    _write_pdf(pdf_path, markdown_path.read_text(encoding="utf-8"))
    return {
        "markdown": str(markdown_path),
        "docx": str(docx_path),
        "pdf": str(pdf_path),
    }


def _write_docx(path: Path, markdown_text: str) -> None:
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10.5)
    for line in markdown_text.splitlines():
        if not line.strip():
            document.add_paragraph("")
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("| "):
            document.add_paragraph(line)
        else:
            document.add_paragraph(line)
    document.save(path)


def _write_pdf(path: Path, markdown_text: str) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    left_margin = 36
    top = height - 36
    y = top
    for line in markdown_text.splitlines():
        if not line.strip():
            y -= 12
        else:
            text = _normalize_pdf_text(line)
            if y < 48:
                pdf.showPage()
                y = top
            pdf.drawString(left_margin, y, text[:120])
            y -= 12
    pdf.save()


def _normalize_pdf_text(text: str) -> str:
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text.encode("latin-1", errors="replace").decode("latin-1")
