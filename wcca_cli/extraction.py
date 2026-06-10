from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .io import write_yaml


HEADER_ALIASES = {
    "ref": "reference",
    "refs": "reference",
    "reference": "reference",
    "reference designator": "reference",
    "designator": "reference",
    "part": "part_number",
    "part number": "part_number",
    "part_number": "part_number",
    "mpn": "part_number",
    "manufacturer part number": "part_number",
    "value": "value",
    "val": "value",
    "description": "description",
    "desc": "description",
    "tolerance": "tolerance",
    "tol": "tolerance",
    "package": "package",
    "footprint": "package",
    "manufacturer": "manufacturer",
    "mfg": "manufacturer",
    "quantity": "quantity",
    "qty": "quantity",
    "pin": "pin",
    "pin number": "pin",
    "pin_number": "pin",
    "net": "net",
    "net name": "net",
    "net_name": "net",
    "requirement id": "requirement_id",
    "requirement_id": "requirement_id",
    "req id": "requirement_id",
    "req_id": "requirement_id",
    "model id": "model_id",
    "model_id": "model_id",
    "circuit type": "model_id",
    "circuit_type": "model_id",
    "min": "min_value",
    "minimum": "min_value",
    "min value": "min_value",
    "min_value": "min_value",
    "max": "max_value",
    "maximum": "max_value",
    "max value": "max_value",
    "max_value": "max_value",
    "unit": "unit",
}


def extract_project(project_dir: str | Path) -> dict[str, Any]:
    root = Path(project_dir)
    components: list[dict[str, Any]] = []
    for path in sorted((root / "input" / "bom").glob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".csv":
            rows = _read_csv(path)
        elif suffix in {".xlsx", ".xlsm"}:
            rows = _read_xlsx(path)
        else:
            continue
        for row_index, row in enumerate(rows, start=2):
            normalized = _normalize_row(row)
            component = _component_from_row(normalized, source_path=path, row_index=row_index, root=root)
            if component:
                components.append(component)
    netlist = _extract_netlist(root)
    requirements = _extract_requirements(root)
    output = {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "components": components,
        "summary": {
            "component_count": len(components),
            "resistor_count": sum(1 for item in components if item.get("component_type") == "resistor"),
            "capacitor_count": sum(1 for item in components if item.get("component_type") == "capacitor"),
        },
    }
    write_yaml(root / "work" / "extracted" / "components.yaml", output)
    write_yaml(root / "work" / "extracted" / "parameter_draft.yaml", _parameter_draft(components))
    write_yaml(root / "work" / "extracted" / "netlist.yaml", netlist)
    write_yaml(root / "work" / "extracted" / "requirements.yaml", requirements)
    return output


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = ["" if value is None else str(value) for value in rows[0]]
    records = []
    for row in rows[1:]:
        records.append({headers[index]: value for index, value in enumerate(row) if index < len(headers)})
    return records


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in row.items():
        canonical = HEADER_ALIASES.get(str(key).strip().lower(), str(key).strip().lower())
        normalized[canonical] = "" if value is None else str(value).strip()
    return normalized


def _component_from_row(row: dict[str, Any], source_path: Path, row_index: int, root: Path) -> dict[str, Any] | None:
    reference = row.get("reference", "").strip()
    if not reference:
        return None
    value_text = row.get("value", "")
    description = row.get("description", "")
    component_type = _component_type(reference, description, value_text)
    component = {
        "reference": reference,
        "component_type": component_type,
        "part_number": row.get("part_number", ""),
        "manufacturer": row.get("manufacturer", ""),
        "package": row.get("package", ""),
        "description": description,
        "source": {
            "type": "bom",
            "path": str(source_path.relative_to(root)).replace("\\", "/"),
            "row": row_index,
        },
        "review_status": "rule_extracted",
    }
    nominal = _parse_nominal_value(value_text, component_type)
    if nominal:
        component["nominal"] = {
            "value": nominal["value"],
            "unit": nominal["unit"],
            "raw": value_text,
            "parameter_type": "nominal",
            "review_status": "rule_extracted",
        }
    tolerance = _parse_tolerance(row.get("tolerance", "") or value_text)
    if tolerance is not None:
        component["tolerance"] = {
            "value": tolerance,
            "unit": "%",
            "parameter_type": "tolerance",
            "review_status": "rule_extracted",
        }
    return component


def _component_type(reference: str, description: str, value_text: str) -> str:
    prefix = re.match(r"[A-Za-z]+", reference)
    prefix_text = prefix.group(0).upper() if prefix else ""
    haystack = f"{description} {value_text}".lower()
    if prefix_text.startswith("R") or "resistor" in haystack or "ohm" in haystack:
        return "resistor"
    if prefix_text.startswith("C") or "capacitor" in haystack or "farad" in haystack:
        return "capacitor"
    if prefix_text.startswith("U"):
        return "ic"
    if prefix_text.startswith("L"):
        return "inductor"
    if prefix_text.startswith("D"):
        return "diode"
    if prefix_text.startswith("Q"):
        return "transistor"
    return "unknown"


def _parse_nominal_value(value_text: str, component_type: str) -> dict[str, Any] | None:
    text = value_text.strip()
    if not text:
        return None
    if component_type == "resistor":
        parsed = _parse_resistance(text)
        if parsed is not None:
            return {"value": parsed, "unit": "ohm"}
    if component_type == "capacitor":
        parsed = _parse_capacitance(text)
        if parsed is not None:
            return {"value": parsed, "unit": "F"}
    return None


def _parse_resistance(text: str) -> float | None:
    compact = text.replace("\u03a9", "ohm").replace(" ", "").lower()
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)(meg|mohm|kohm|k|ohm|r)?", compact)
    if not match:
        return None
    value = float(match.group(1))
    suffix = match.group(2) or "ohm"
    if suffix in {"meg", "mohm"}:
        return value * 1_000_000.0
    if suffix in {"k", "kohm"}:
        return value * 1_000.0
    return value


def _parse_capacitance(text: str) -> float | None:
    compact = text.replace("\u00b5", "u").replace("\u03bc", "u").replace(" ", "").lower()
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)(pf|nf|uf|mf|f)", compact)
    if not match:
        return None
    value = float(match.group(1))
    suffix = match.group(2)
    factors = {"pf": 1e-12, "nf": 1e-9, "uf": 1e-6, "mf": 1e-3, "f": 1.0}
    return value * factors[suffix]


def _parse_tolerance(text: str) -> float | None:
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*%", text)
    if not match:
        return None
    return float(match.group(1))


def _parameter_draft(components: list[dict[str, Any]]) -> dict[str, Any]:
    parameters = []
    for component in components:
        for field in ("nominal", "tolerance"):
            record = component.get(field)
            if not record:
                continue
            parameters.append({
                "component_ref": component["reference"],
                "component_type": component["component_type"],
                "parameter": field,
                "value": record["value"],
                "unit": record["unit"],
                "source": component["source"],
                "review_status": "rule_extracted",
            })
    return {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "parameters": parameters,
    }


def _extract_netlist(root: Path) -> dict[str, Any]:
    connections: list[dict[str, Any]] = []
    for path in sorted((root / "input" / "netlist").glob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".csv":
            rows = _read_csv(path)
            connections.extend(_connections_from_rows(rows, path, root))
        elif suffix in {".net", ".txt"}:
            connections.extend(_connections_from_text(path, root))
    nets: dict[str, list[str]] = {}
    component_pins: dict[str, dict[str, str]] = {}
    for connection in connections:
        nets.setdefault(connection["net"], []).append(f"{connection['reference']}.{connection['pin']}")
        component_pins.setdefault(connection["reference"], {})[connection["pin"]] = connection["net"]
    return {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "connections": connections,
        "nets": [{"name": net, "nodes": sorted(nodes)} for net, nodes in sorted(nets.items())],
        "component_pins": component_pins,
        "topology_hints": _topology_hints(component_pins),
        "summary": {
            "connection_count": len(connections),
            "net_count": len(nets),
            "component_count": len(component_pins),
        },
    }


def _connections_from_rows(rows: list[dict[str, Any]], path: Path, root: Path) -> list[dict[str, Any]]:
    connections = []
    for row_index, row in enumerate(rows, start=2):
        normalized = _normalize_row(row)
        reference = normalized.get("reference", "")
        pin = normalized.get("pin", normalized.get("pin_number", ""))
        net = normalized.get("net", normalized.get("net_name", ""))
        if reference and pin and net:
            connections.append(_connection(reference, pin, net, path, root, row_index))
    return connections


def _connections_from_text(path: Path, root: Path) -> list[dict[str, Any]]:
    connections = []
    for line_index, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        parts = re.split(r"[\s,]+", stripped)
        if len(parts) >= 3:
            reference, pin, net = parts[:3]
            connections.append(_connection(reference, pin, net, path, root, line_index))
    return connections


def _connection(reference: str, pin: str, net: str, path: Path, root: Path, row: int) -> dict[str, Any]:
    return {
        "reference": reference.strip(),
        "pin": str(pin).strip(),
        "net": net.strip(),
        "source": {
            "type": "netlist",
            "path": str(path.relative_to(root)).replace("\\", "/"),
            "row": row,
        },
        "review_status": "rule_extracted",
    }


def _topology_hints(component_pins: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    hints: list[dict[str, Any]] = []
    refs = sorted(component_pins)
    for first in refs:
        for second in refs:
            if first >= second:
                continue
            first_pins = component_pins[first]
            second_pins = component_pins[second]
            common = sorted(
                net for net in set(first_pins.values()) & set(second_pins.values())
                if not _is_global_net(net)
            )
            if not common:
                continue
            if first.startswith("R") and second.startswith("R"):
                candidate_type = "resistor_pair"
            elif (first.startswith("R") and second.startswith("C")) or (first.startswith("C") and second.startswith("R")):
                candidate_type = "rc_pair"
            else:
                candidate_type = "shared_net"
            hints.append({
                "candidate_type": candidate_type,
                "component_refs": [first, second],
                "shared_nets": common,
                "review_status": "rule_extracted",
            })
    return hints


def _is_global_net(net: str) -> bool:
    normalized = net.strip().upper()
    if normalized in {"GND", "PGND", "AGND", "DGND", "0", "3V3", "5V", "VBAT", "VBAT_5V", "VCC", "VDD"}:
        return True
    return normalized.startswith("VCC") or normalized.startswith("VDD")


def _extract_requirements(root: Path) -> dict[str, Any]:
    requirements: list[dict[str, Any]] = []
    for path in sorted((root / "input" / "requirements").glob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".csv":
            rows = _read_csv(path)
        elif suffix in {".xlsx", ".xlsm"}:
            rows = _read_xlsx(path)
        else:
            continue
        for row_index, row in enumerate(rows, start=2):
            normalized = _normalize_row(row)
            requirement = _requirement_from_row(normalized, path, root, row_index)
            if requirement:
                requirements.append(requirement)
    return {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "requirements": requirements,
        "summary": {
            "requirement_count": len(requirements),
            "mapped_requirement_count": sum(1 for item in requirements if item.get("model_id")),
        },
    }


def _requirement_from_row(row: dict[str, Any], path: Path, root: Path, row_index: int) -> dict[str, Any] | None:
    requirement_id = row.get("requirement_id", "").strip()
    if not requirement_id:
        return None
    return {
        "id": requirement_id,
        "model_id": row.get("model_id", "").strip(),
        "description": row.get("description", "").strip() or "Requirement imported from requirements table.",
        "min_value": _parse_number(row.get("min_value", "")),
        "max_value": _parse_number(row.get("max_value", "")),
        "unit": row.get("unit", "").strip() or None,
        "source": {
            "type": "requirement",
            "path": str(path.relative_to(root)).replace("\\", "/"),
            "row": row_index,
        },
        "review_status": "rule_extracted",
    }


def _parse_number(value: Any) -> float | None:
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None
